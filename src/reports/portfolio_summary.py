"""Module providing N100 financial intelligence functionality."""

import csv
import math
import sqlite3
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parents[2]
DB_PATH = ROOT / "data" / "db" / "n100.db"

OUTPUT_DIR = ROOT / "reports" / "portfolio"
OUTPUT_FILE = OUTPUT_DIR / "N100_Portfolio_Summary.pdf"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# PAGE SETTINGS
# ============================================================

PAGE_W, PAGE_H = A4
MARGIN = 14 * mm
CONTENT_W = PAGE_W - 2 * MARGIN


# ============================================================
# HELPERS
# ============================================================


def safe_float(value):
    """Handle safe float."""
    try:
        if value is None:
            return None

        value = float(value)

        if math.isnan(value) or math.isinf(value):
            return None

        return value

    except (TypeError, ValueError):
        return None


def fmt(value, decimals=2):
    """Handle fmt."""
    value = safe_float(value)

    if value is None:
        return "N/A"

    return f"{value:,.{decimals}f}"


def fmt_pct(value):
    """Handle fmt pct."""
    value = safe_float(value)

    if value is None:
        return "N/A"

    return f"{value:.1f}%"


def trend_arrow(current, previous):
    """Handle trend arrow."""
    current = safe_float(current)
    previous = safe_float(previous)

    if current is None or previous is None:
        return "→"

    if current > previous:
        return "↑"

    if current < previous:
        return "↓"

    return "→"


def latest_by_company(rows):
    """Handle latest by company."""
    result = {}

    for row in rows:
        result[row["company_id"]] = row

    return result


def previous_by_company(rows):
    """Handle previous by company."""
    result = {}
    grouped = {}

    for row in rows:
        grouped.setdefault(row["company_id"], []).append(row)

    for company_id, company_rows in grouped.items():
        if len(company_rows) >= 2:
            result[company_id] = company_rows[-2]

    return result


# ============================================================
# DATABASE
# ============================================================


def load_data():
    """Retrieve data."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    companies = conn.execute("""
        SELECT
            c.id AS company_id,
            c.company_name,
            s.broad_sector,
            s.sub_sector
        FROM companies c
        LEFT JOIN sectors s
            ON s.company_id = c.id
        ORDER BY c.id
        """).fetchall()

    ratios = conn.execute("""
        SELECT *
        FROM financial_ratios
        ORDER BY company_id, year
        """).fetchall()

    pl = conn.execute("""
        SELECT *
        FROM profitandloss
        ORDER BY company_id, year
        """).fetchall()

    cf = conn.execute("""
        SELECT *
        FROM cashflow
        ORDER BY company_id, year
        """).fetchall()

    conn.close()

    return companies, ratios, pl, cf


# ============================================================
# OUTPUT DATA
# ============================================================


def load_cashflow_intelligence():
    """Retrieve cashflow intelligence."""
    path = ROOT / "output" / "cashflow_intelligence.xlsx"

    if not path.exists():
        return {}

    try:
        import openpyxl

        wb = openpyxl.load_workbook(
            path,
            data_only=True,
        )

        ws = wb.active

        headers = [cell.value for cell in ws[1]]

        result = {}

        for values in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):
            row = dict(zip(headers, values))

            company_id = row.get("company_id")

            if company_id is not None:
                result[str(company_id)] = row

        return result

    except Exception:
        return {}


def load_capital_allocation():
    """Retrieve capital allocation."""
    path = ROOT / "output" / "capital_allocation.csv"

    if not path.exists():
        return {}

    result = {}

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.DictReader(f)

        for row in reader:
            company_id = str(row.get("company_id", "")).strip()

            if not company_id:
                continue

            result[company_id] = row

    return result


# ============================================================
# DRAWING
# ============================================================


def draw_header(c, title, subtitle):
    """Render header."""
    c.setFillColor(colors.HexColor("#17365D"))

    c.rect(
        0,
        PAGE_H - 19 * mm,
        PAGE_W,
        19 * mm,
        fill=1,
        stroke=0,
    )

    c.setFillColor(colors.white)

    c.setFont(
        "Helvetica-Bold",
        14,
    )

    c.drawString(
        MARGIN,
        PAGE_H - 11.5 * mm,
        title[:75],
    )

    c.setFont(
        "Helvetica",
        7.5,
    )

    c.drawRightString(
        PAGE_W - MARGIN,
        PAGE_H - 11.5 * mm,
        subtitle[:75],
    )


def draw_footer(c, page_number, total_pages):
    """Render footer."""
    c.setStrokeColor(colors.HexColor("#C8D2DC"))

    c.line(
        MARGIN,
        10 * mm,
        PAGE_W - MARGIN,
        10 * mm,
    )

    c.setFillColor(colors.HexColor("#666666"))

    c.setFont(
        "Helvetica",
        7,
    )

    c.drawString(
        MARGIN,
        6.5 * mm,
        "N100 Financial Intelligence Platform",
    )

    c.drawRightString(
        PAGE_W - MARGIN,
        6.5 * mm,
        f"Page {page_number} of {total_pages}",
    )


def section_title(c, x, y, title, width=CONTENT_W):
    """Handle section title."""
    c.setFillColor(colors.HexColor("#EAF0F5"))

    c.roundRect(
        x,
        y - 5 * mm,
        width,
        8 * mm,
        2 * mm,
        fill=1,
        stroke=0,
    )

    c.setFillColor(colors.HexColor("#17365D"))

    c.setFont(
        "Helvetica-Bold",
        9,
    )

    c.drawString(
        x + 3 * mm,
        y - 2 * mm,
        title,
    )

    return y - 10 * mm


def draw_table(
    c,
    x,
    y,
    widths,
    rows,
    row_h=7 * mm,
    font_size=6.5,
):
    """Render table."""
    if not rows:
        return y

    total_width = sum(widths)

    for row_index, row in enumerate(rows):

        if row_index == 0:
            fill = colors.HexColor("#17365D")
            text_color = colors.white
            font = "Helvetica-Bold"
        else:
            fill = colors.HexColor("#F7F9FB") if row_index % 2 == 0 else colors.white
            text_color = colors.HexColor("#222222")
            font = "Helvetica"

        c.setFillColor(fill)

        c.setStrokeColor(colors.HexColor("#D5DDE4"))

        c.rect(
            x,
            y - row_h,
            total_width,
            row_h,
            fill=1,
            stroke=1,
        )

        current_x = x

        for i, cell in enumerate(row):

            text = "" if cell is None else str(cell)

            text = text.replace("\n", " ")[:34]

            c.setFillColor(text_color)

            c.setFont(
                font,
                font_size,
            )

            c.drawString(
                current_x + 2 * mm,
                y - row_h + 2.3 * mm,
                text,
            )

            current_x += widths[i]

        y -= row_h

    return y


def draw_kpi_card(
    c,
    x,
    y,
    width,
    height,
    label,
    value,
):
    """Render kpi card."""
    c.setFillColor(colors.white)

    c.setStrokeColor(colors.HexColor("#D0D7DE"))

    c.roundRect(
        x,
        y - height,
        width,
        height,
        2 * mm,
        fill=1,
        stroke=1,
    )

    c.setFillColor(colors.HexColor("#666666"))

    c.setFont(
        "Helvetica",
        7,
    )

    c.drawString(
        x + 3 * mm,
        y - 5 * mm,
        label[:24],
    )

    c.setFillColor(colors.HexColor("#17365D"))

    c.setFont(
        "Helvetica-Bold",
        11,
    )

    c.drawString(
        x + 3 * mm,
        y - 11 * mm,
        str(value)[:20],
    )


# ============================================================
# PORTFOLIO OVERVIEW
# ============================================================


def draw_overview_page(
    c,
    companies,
    ratios,
    pl,
    cf,
    intelligence,
    allocations,
    total_pages,
):
    """Render overview page."""
    draw_header(
        c,
        "N100 FINANCIAL INTELLIGENCE",
        "Portfolio Summary",
    )

    y = PAGE_H - 29 * mm

    y = section_title(
        c,
        MARGIN,
        y,
        "PORTFOLIO OVERVIEW",
    )

    company_count = len(companies)

    sector_counts = {}

    for company in companies:
        sector = company["broad_sector"] or "Unclassified"

        sector_counts[sector] = sector_counts.get(sector, 0) + 1

    distress_counts = {
        "High": 0,
        "Medium": 0,
        "Low": 0,
    }

    for row in intelligence.values():

        level = str(row.get("distress_level", "")).strip()

        if level in distress_counts:
            distress_counts[level] += 1

    allocation_counts = {}

    for row in allocations.values():

        pattern = str(row.get("pattern_label", "N/A")).strip()

        allocation_counts[pattern] = allocation_counts.get(pattern, 0) + 1

    # KPI cards
    gap = 3 * mm
    card_w = (CONTENT_W - 3 * gap) / 4

    card_h = 19 * mm

    cards = [
        (
            "Companies",
            company_count,
        ),
        (
            "Sectors",
            len(sector_counts),
        ),
        (
            "High Distress",
            distress_counts["High"],
        ),
        (
            "Medium Distress",
            distress_counts["Medium"],
        ),
    ]

    x = MARGIN

    for label, value in cards:

        draw_kpi_card(
            c,
            x,
            y,
            card_w,
            card_h,
            label,
            value,
        )

        x += card_w + gap

    y -= 25 * mm

    # Sector distribution
    y = section_title(
        c,
        MARGIN,
        y,
        "SECTOR DISTRIBUTION",
    )

    sector_rows = [
        [
            "Sector",
            "Companies",
            "Portfolio %",
        ]
    ]

    for sector, count in sorted(
        sector_counts.items(),
        key=lambda item: (-item[1], item[0]),
    ):

        pct = count / company_count * 100 if company_count else 0

        sector_rows.append(
            [
                sector,
                count,
                f"{pct:.1f}%",
            ]
        )

    # Keep overview compact.
    sector_rows = sector_rows[:13]

    y = draw_table(
        c,
        MARGIN,
        y,
        [
            78 * mm,
            35 * mm,
            35 * mm,
        ],
        sector_rows,
        row_h=6 * mm,
        font_size=6.2,
    )

    y -= 7 * mm

    # Risk and capital allocation
    half_gap = 6 * mm
    half_w = (CONTENT_W - half_gap) / 2

    y_left = section_title(
        c,
        MARGIN,
        y,
        "DISTRESS PROFILE",
        width=half_w,
    )

    distress_rows = [
        ["Level", "Companies"],
        ["High", distress_counts["High"]],
        ["Medium", distress_counts["Medium"]],
        ["Low", distress_counts["Low"]],
    ]

    y_left = draw_table(
        c,
        MARGIN,
        y_left,
        [
            38 * mm,
            35 * mm,
        ],
        distress_rows,
        row_h=6 * mm,
        font_size=6.2,
    )

    x_right = MARGIN + half_w + half_gap

    y_right = section_title(
        c,
        x_right,
        y,
        "CAPITAL ALLOCATION",
        width=half_w,
    )

    allocation_rows = [
        ["Pattern", "Rows"],
    ]

    for pattern, count in sorted(
        allocation_counts.items(),
        key=lambda item: (-item[1], item[0]),
    )[:8]:

        allocation_rows.append(
            [
                pattern,
                count,
            ]
        )

    draw_table(
        c,
        x_right,
        y_right,
        [
            58 * mm,
            20 * mm,
        ],
        allocation_rows,
        row_h=5.8 * mm,
        font_size=5.8,
    )

    # Footer
    draw_footer(
        c,
        1,
        total_pages,
    )


# ============================================================
# COMPANY PAGE
# ============================================================


def draw_company_page(
    c,
    company,
    ratios_by_company,
    pl_by_company,
    cf_by_company,
    intelligence,
    allocations,
    page_number,
    total_pages,
):
    """Render company page."""
    company_id = str(company["company_id"])

    company_name = str(company["company_name"])

    sector = company["broad_sector"] or "Unclassified"

    sub_sector = company["sub_sector"] or "N/A"

    ratios = ratios_by_company.get(
        company_id,
        [],
    )

    pl = pl_by_company.get(
        company_id,
        [],
    )

    cf = cf_by_company.get(
        company_id,
        [],
    )

    latest_ratio = ratios[-1] if ratios else None

    previous_ratio = ratios[-2] if len(ratios) >= 2 else None

    latest_pl = pl[-1] if pl else None

    pl[-2] if len(pl) >= 2 else None

    latest_cf = cf[-1] if cf else None

    draw_header(
        c,
        company_name,
        f"{company_id} | {sector}",
    )

    y = PAGE_H - 28 * mm

    c.setFillColor(colors.HexColor("#444444"))

    c.setFont(
        "Helvetica",
        7.5,
    )

    c.drawString(
        MARGIN,
        y,
        f"Sector: {sector}    |    Sub-sector: {sub_sector}",
    )

    y -= 8 * mm

    # Company KPIs
    gap = 3 * mm
    card_w = (CONTENT_W - 3 * gap) / 4

    card_h = 18 * mm

    cards = [
        (
            "Revenue",
            fmt(latest_pl["sales"] if latest_pl else None),
        ),
        (
            "Net Profit",
            fmt(latest_pl["net_profit"] if latest_pl else None),
        ),
        (
            "CFO",
            fmt(latest_cf["operating_activity"] if latest_cf else None),
        ),
        (
            "FCF",
            fmt(latest_ratio["free_cash_flow_cr"] if latest_ratio else None),
        ),
    ]

    x = MARGIN

    for label, value in cards:

        draw_kpi_card(
            c,
            x,
            y,
            card_w,
            card_h,
            label,
            value,
        )

        x += card_w + gap

    y -= 24 * mm

    # KPI trend table
    y = section_title(
        c,
        MARGIN,
        y,
        "KEY PERFORMANCE INDICATORS",
    )

    kpi_rows = [
        [
            "Metric",
            "Latest",
            "Previous",
            "Trend",
        ],
        [
            "ROE",
            fmt_pct(latest_ratio["return_on_equity_pct"] if latest_ratio else None),
            fmt_pct(previous_ratio["return_on_equity_pct"] if previous_ratio else None),
            trend_arrow(
                latest_ratio["return_on_equity_pct"] if latest_ratio else None,
                previous_ratio["return_on_equity_pct"] if previous_ratio else None,
            ),
        ],
        [
            "Revenue CAGR 5Y",
            fmt_pct(latest_ratio["revenue_cagr_5yr"] if latest_ratio else None),
            "—",
            "→",
        ],
        [
            "PAT CAGR 5Y",
            fmt_pct(latest_ratio["pat_cagr_5yr"] if latest_ratio else None),
            "—",
            "→",
        ],
        [
            "EPS CAGR 5Y",
            fmt_pct(latest_ratio["eps_cagr_5yr"] if latest_ratio else None),
            "—",
            "→",
        ],
        [
            "Debt / Equity",
            fmt(latest_ratio["debt_to_equity"] if latest_ratio else None),
            fmt(previous_ratio["debt_to_equity"] if previous_ratio else None),
            trend_arrow(
                previous_ratio["debt_to_equity"] if previous_ratio else None,
                latest_ratio["debt_to_equity"] if latest_ratio else None,
            ),
        ],
        [
            "Interest Coverage",
            fmt(latest_ratio["interest_coverage"] if latest_ratio else None),
            fmt(previous_ratio["interest_coverage"] if previous_ratio else None),
            trend_arrow(
                latest_ratio["interest_coverage"] if latest_ratio else None,
                previous_ratio["interest_coverage"] if previous_ratio else None,
            ),
        ],
    ]

    y = draw_table(
        c,
        MARGIN,
        y,
        [
            62 * mm,
            35 * mm,
            35 * mm,
            25 * mm,
        ],
        kpi_rows,
        row_h=6.5 * mm,
        font_size=6.2,
    )

    y -= 7 * mm

    # Financial history
    y = section_title(
        c,
        MARGIN,
        y,
        "RECENT FINANCIAL HISTORY",
    )

    history_rows = [
        [
            "Year",
            "Sales",
            "PAT",
            "CFO",
            "FCF",
        ]
    ]

    # Map ratios/cashflow by year
    ratio_by_year = {str(row["year"]): row for row in ratios}

    cf_by_year = {str(row["year"]): row for row in cf}

    for pl_row in pl[-6:]:

        year = str(pl_row["year"])

        ratio = ratio_by_year.get(year)
        cf_row = cf_by_year.get(year)

        history_rows.append(
            [
                year,
                fmt(pl_row["sales"]),
                fmt(pl_row["net_profit"]),
                fmt(cf_row["operating_activity"] if cf_row else None),
                fmt(ratio["free_cash_flow_cr"] if ratio else None),
            ]
        )

    y = draw_table(
        c,
        MARGIN,
        y,
        [
            27 * mm,
            32 * mm,
            32 * mm,
            32 * mm,
            32 * mm,
        ],
        history_rows,
        row_h=6 * mm,
        font_size=5.9,
    )

    y -= 7 * mm

    # Intelligence
    y = section_title(
        c,
        MARGIN,
        y,
        "CASH FLOW & CAPITAL ALLOCATION",
    )

    intelligence_row = intelligence.get(
        company_id,
        {},
    )

    allocation_row = allocations.get(
        company_id,
        {},
    )

    intelligence_rows = [
        [
            "Signal",
            "Assessment",
        ],
        [
            "CFO / PAT Quality",
            str(
                intelligence_row.get(
                    "cfo_pat_quality",
                    "N/A",
                )
            ),
        ],
        [
            "CapEx Intensity",
            str(
                intelligence_row.get(
                    "capex_intensity",
                    "N/A",
                )
            ),
        ],
        [
            "Distress Level",
            str(
                intelligence_row.get(
                    "distress_level",
                    "N/A",
                )
            ),
        ],
        [
            "Capital Allocation",
            str(
                allocation_row.get(
                    "pattern_label",
                    "N/A",
                )
            ),
        ],
    ]

    draw_table(
        c,
        MARGIN,
        y,
        [
            62 * mm,
            93 * mm,
        ],
        intelligence_rows,
        row_h=6.5 * mm,
        font_size=6.1,
    )

    draw_footer(
        c,
        page_number,
        total_pages,
    )


# ============================================================
# MAIN
# ============================================================


def main():
    """Run the module's main workflow."""
    print("=== DAY 35 PORTFOLIO SUMMARY ===")

    (
        companies,
        ratios,
        pl,
        cf,
    ) = load_data()

    print(f"Companies found: {len(companies)}")

    intelligence = load_cashflow_intelligence()
    allocations = load_capital_allocation()

    ratios_by_company = {}
    pl_by_company = {}
    cf_by_company = {}

    for row in ratios:
        ratios_by_company.setdefault(
            str(row["company_id"]),
            [],
        ).append(row)

    for row in pl:
        pl_by_company.setdefault(
            str(row["company_id"]),
            [],
        ).append(row)

    for row in cf:
        cf_by_company.setdefault(
            str(row["company_id"]),
            [],
        ).append(row)

    total_pages = 1 + len(companies)

    c = canvas.Canvas(
        str(OUTPUT_FILE),
        pagesize=A4,
        pageCompression=1,
    )

    c.setTitle("N100 Financial Intelligence - Portfolio Summary")

    c.setAuthor("N100 Financial Intelligence Platform")

    # --------------------------------------------------------
    # Page 1: portfolio overview
    # --------------------------------------------------------

    draw_overview_page(
        c,
        companies,
        ratios,
        pl,
        cf,
        intelligence,
        allocations,
        total_pages,
    )

    c.showPage()

    # --------------------------------------------------------
    # One page per company
    # --------------------------------------------------------

    for index, company in enumerate(
        companies,
        start=2,
    ):

        company_id = str(company["company_id"])

        draw_company_page(
            c,
            company,
            ratios_by_company,
            pl_by_company,
            cf_by_company,
            intelligence,
            allocations,
            index,
            total_pages,
        )

        c.showPage()

        print(f"Added page: {company_id} | " f"{company['company_name']}")

    c.save()

    print()
    print(f"Portfolio pages: {total_pages}")

    print(f"Output: {OUTPUT_FILE}")

    print("=== DAY 35 PORTFOLIO SUMMARY COMPLETE ===")


if __name__ == "__main__":
    main()
