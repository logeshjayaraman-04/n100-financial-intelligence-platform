"""
Sprint 3 - Day 17
Composite Quality Score and Screener Excel Export.
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from .engine import build_screener_dataframe
from .presets import PRESETS, run_preset, run_turnaround_watch


DB_PATH = Path("data/db/n100.db")
OUTPUT_PATH = Path("output/screener_output.xlsx")


# ============================================================
# HELPERS
# ============================================================

def winsorised_score(series, higher_is_better=True):
    values = pd.to_numeric(series, errors="coerce")
    valid = values.dropna()

    if valid.empty:
        return pd.Series(50.0, index=series.index)

    p10 = valid.quantile(0.10)
    p90 = valid.quantile(0.90)

    if pd.isna(p10) or pd.isna(p90) or p10 == p90:
        return pd.Series(50.0, index=series.index)

    clipped = values.clip(lower=p10, upper=p90)

    score = ((clipped - p10) / (p90 - p10)) * 100

    if not higher_is_better:
        score = 100 - score

    return score.fillna(50.0)


def find_column(df, candidates):
    for column in candidates:
        if column in df.columns:
            return column
    return None


# ============================================================
# COMPOSITE SCORE
# ============================================================

def calculate_composite_score(df):
    result = df.copy()

    roe_col = find_column(
        result,
        ["roe", "return_on_equity_pct"],
    )

    roce_col = find_column(
        result,
        ["roce", "return_on_capital_employed_pct"],
    )

    npm_col = find_column(
        result,
        ["npm", "net_profit_margin_pct"],
    )

    fcf_col = find_column(
        result,
        ["fcf", "free_cash_flow_cr"],
    )

    cfo_col = find_column(
        result,
        ["cfo", "cash_from_operations_cr"],
    )

    pat_col = find_column(
        result,
        ["pat", "net_profit"],
    )

    revenue_cagr_col = find_column(
        result,
        ["revenue_cagr_5yr"],
    )

    pat_cagr_col = find_column(
        result,
        ["pat_cagr_5yr"],
    )

    de_col = find_column(
        result,
        ["de", "debt_to_equity"],
    )

    icr_col = find_column(
        result,
        ["icr", "interest_coverage"],
    )

    # --------------------------------------------------------
    # Profitability: 35%
    # --------------------------------------------------------

    if roe_col:
        roe_score = winsorised_score(
            result[roe_col],
            True,
        )
    else:
        roe_score = pd.Series(
            50.0,
            index=result.index,
        )

    if roce_col:
        roce_score = winsorised_score(
            result[roce_col],
            True,
        )
    else:
        roce_score = pd.Series(
            50.0,
            index=result.index,
        )

    if npm_col:
        npm_score = winsorised_score(
            result[npm_col],
            True,
        )
    else:
        npm_score = pd.Series(
            50.0,
            index=result.index,
        )

    profitability_score = (
        roe_score * 0.15
        + roce_score * 0.10
        + npm_score * 0.10
    )

    # --------------------------------------------------------
    # Cash Quality: 30%
    # --------------------------------------------------------

    if fcf_col:
        fcf_score = winsorised_score(
            result[fcf_col],
            True,
        )

        fcf_positive = (
            pd.to_numeric(
                result[fcf_col],
                errors="coerce",
            ) > 0
        ).astype(float) * 100

    else:
        fcf_score = pd.Series(
            50.0,
            index=result.index,
        )

        fcf_positive = pd.Series(
            50.0,
            index=result.index,
        )

    if cfo_col and pat_col:

        cfo_values = pd.to_numeric(
            result[cfo_col],
            errors="coerce",
        )

        pat_values = pd.to_numeric(
            result[pat_col],
            errors="coerce",
        )

        cfo_pat_ratio = pd.Series(
            np.nan,
            index=result.index,
        )

        valid = (
            pat_values.notna()
            & (pat_values != 0)
        )

        cfo_pat_ratio.loc[valid] = (
            cfo_values.loc[valid]
            / pat_values.loc[valid]
        )

        cfo_pat_score = winsorised_score(
            cfo_pat_ratio,
            True,
        )

        result["cfo_pat_ratio"] = cfo_pat_ratio

    else:

        cfo_pat_score = pd.Series(
            50.0,
            index=result.index,
        )

        result["cfo_pat_ratio"] = np.nan

    cash_quality_score = (
        fcf_score * 0.15
        + cfo_pat_score * 0.10
        + fcf_positive * 0.05
    )

    # --------------------------------------------------------
    # Growth: 20%
    # --------------------------------------------------------

    if revenue_cagr_col:
        revenue_growth_score = winsorised_score(
            result[revenue_cagr_col],
            True,
        )
    else:
        revenue_growth_score = pd.Series(
            50.0,
            index=result.index,
        )

    if pat_cagr_col:
        pat_growth_score = winsorised_score(
            result[pat_cagr_col],
            True,
        )
    else:
        pat_growth_score = pd.Series(
            50.0,
            index=result.index,
        )

    growth_score = (
        revenue_growth_score * 0.10
        + pat_growth_score * 0.10
    )

    # --------------------------------------------------------
    # Leverage: 15%
    # --------------------------------------------------------

    if de_col:
        de_score = winsorised_score(
            result[de_col],
            higher_is_better=False,
        )
    else:
        de_score = pd.Series(
            50.0,
            index=result.index,
        )

    if icr_col:

        icr_values = pd.to_numeric(
            result[icr_col],
            errors="coerce",
        )

        # Debt-free companies have no interest burden.
        # Treat them as infinite ICR.

        icr_values = icr_values.fillna(np.inf)

        icr_score = winsorised_score(
            icr_values,
            True,
        )

    else:
        icr_score = pd.Series(
            50.0,
            index=result.index,
        )

    leverage_score = (
        de_score * 0.10
        + icr_score * 0.05
    )

    # --------------------------------------------------------
    # Final 0-100 score
    # --------------------------------------------------------

    result["composite_quality_score"] = (
        profitability_score
        + cash_quality_score
        + growth_score
        + leverage_score
    )

    result["composite_quality_score"] = (
        result["composite_quality_score"]
        .clip(0, 100)
        .round(2)
    )

    return result


# ============================================================
# PRESET DATA
# ============================================================

def get_preset_dataframe(preset_name, db_path=DB_PATH):

    if preset_name == "Turnaround Watch":

        df = run_turnaround_watch(
            db_path
        )

    else:

        df = run_preset(
            preset_name,
            db_path,
        )

    return calculate_composite_score(df)


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_screener_excel(
    output_path=OUTPUT_PATH,
    db_path=DB_PATH,
):

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    preset_data = {}

    preset_names = list(PRESETS)

    if "Turnaround Watch" not in preset_names:
        preset_names.append(
            "Turnaround Watch"
        )

    for preset_name in preset_names:

        print(
            f"Preparing: {preset_name}"
        )

        preset_data[preset_name] = (
            get_preset_dataframe(
                preset_name,
                db_path,
            )
        )

    # --------------------------------------------------------
    # Write Excel
    # --------------------------------------------------------

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        for preset_name, df in preset_data.items():

            export_df = df.copy()

            if "composite_quality_score" in export_df.columns:

                export_df = export_df.sort_values(
                    "composite_quality_score",
                    ascending=False,
                )

            sheet_name = preset_name[:31]

            export_df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    # --------------------------------------------------------
    # Format workbook
    # --------------------------------------------------------

    workbook = load_workbook(
        output_path
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    for worksheet in workbook.worksheets:

        # Header

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = header_fill

        # Freeze header

        worksheet.freeze_panes = "A2"

        # Filter

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        # Column widths

        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 10),
                30,
            )

        # ----------------------------------------------------
        # Composite score colour coding
        # ----------------------------------------------------

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        if (
            "composite_quality_score"
            in headers
        ):

            score_column = (
                headers.index(
                    "composite_quality_score"
                ) + 1
            )

            for row in range(
                2,
                worksheet.max_row + 1,
            ):

                cell = worksheet.cell(
                    row=row,
                    column=score_column,
                )

                if isinstance(
                    cell.value,
                    (int, float),
                ):

                    if cell.value >= 75:

                        cell.fill = green_fill

                    elif cell.value < 40:

                        cell.fill = red_fill

    workbook.save(
        output_path
    )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print()
    print("=" * 80)
    print("DAY 17 — SCREENER EXPORT COMPLETE")
    print("=" * 80)
    print(
        f"Output: {output_path}"
    )

    for name, df in preset_data.items():

        if "company_id" in df.columns:

            company_count = (
                df["company_id"]
                .nunique()
            )

        else:

            company_count = 0

        print(
            f"{name}: "
            f"{len(df)} rows / "
            f"{company_count} companies"
        )


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    export_screener_excel()